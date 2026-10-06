import os
import re
import time
import requests
import pandas as pd

from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


BASE_URL = "https://www.technoedgels.com/blog-insights/"

CHECKPOINT_FILE = "blog_checkpoint.csv"
OUTPUT_FILE = "technoedge_all_blogs.xlsx"

REQUEST_DELAY = 1
TIMEOUT = 30


def clean_text(text):

    if not text:
        return ""

    return re.sub(
        r"\s+",
        " ",
        text
    ).strip()


# --------------------------------------------------
# CREATE HTTP SESSION
# --------------------------------------------------

def create_session():

    session = requests.Session()

    session.headers.update({

        "User-Agent": (
            "Mozilla/5.0 "
            "(Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 "
            "(KHTML, like Gecko) "
            "Chrome/154.0 Safari/537.36"
        ),

        "Accept-Language":
            "en-GB,en;q=0.9"
    })

    retry = Retry(

        total=3,

        backoff_factor=1,

        status_forcelist=[
            429,
            500,
            502,
            503,
            504
        ],

        allowed_methods=[
            "GET"
        ]
    )

    adapter = HTTPAdapter(
        max_retries=retry
    )

    session.mount(
        "https://",
        adapter
    )

    session.mount(
        "http://",
        adapter
    )

    return session


session = create_session()


# --------------------------------------------------
# DOWNLOAD PAGE
# --------------------------------------------------

def get_soup(url):

    response = session.get(
        url,
        timeout=TIMEOUT
    )

    response.raise_for_status()

    return BeautifulSoup(
        response.text,
        "lxml"
    )


# --------------------------------------------------
# DETECT TOTAL BLOG ARCHIVE PAGES
# --------------------------------------------------

def get_total_pages(soup):

    pagination = soup.select_one(
        ".rt-pagination-wrap[data-total-pages]"
    )

    if pagination:

        total = pagination.get(
            "data-total-pages"
        )

        if total and total.isdigit():

            return int(total)

    return 1


# --------------------------------------------------
# EXTRACT CARDS FROM ARCHIVE PAGE
# --------------------------------------------------

def extract_cards(
    soup,
    page_number
):

    results = []

    cards = soup.select(
        ".rt-grid-item"
    )

    for card in cards:

        title_link = card.select_one(
            "h3.entry-title "
            "a.tpg-post-link"
        )

        if not title_link:
            continue

        date_element = card.select_one(
            ".post-meta-tags .date a"
        )

        excerpt_element = card.select_one(
            ".tpg-excerpt-inner"
        )

        image_element = card.select_one(
            ".rt-img-holder "
            "img.rt-img-responsive"
        )

        blog_url = title_link.get(
            "href",
            ""
        )

        thumbnail_url = ""

        if image_element:

            thumbnail_url = (
                image_element.get("src")
                or
                image_element.get("data-src")
                or
                ""
            )

        results.append({

            "Archive_Page":
                page_number,

            "Post_ID":
                card.get(
                    "data-id",
                    ""
                ),

            "Card_Title":
                clean_text(
                    title_link.get_text()
                ),

            "Card_Date":
                clean_text(
                    date_element.get_text()
                    if date_element
                    else ""
                ),

            "Excerpt":
                clean_text(
                    excerpt_element.get_text()
                    if excerpt_element
                    else ""
                ),

            "Thumbnail_URL":
                thumbnail_url,

            "Blog_URL":
                blog_url
        })

    return results


# --------------------------------------------------
# EXTRACT TABLE
# --------------------------------------------------

def extract_table(table):

    rows = []

    for row in table.select("tr"):

        cells = [

            clean_text(
                cell.get_text(
                    " ",
                    strip=True
                )
            )

            for cell in row.select(
                "th, td"
            )
        ]

        if cells:

            rows.append(
                " | ".join(cells)
            )

    return "\n".join(rows)


# --------------------------------------------------
# EXTRACT ARTICLE BODY IN READING ORDER
# --------------------------------------------------

def extract_article_body(body):

    content = []

    elements = body.find_all(
        [
            "h2",
            "h3",
            "p",
            "li",
            "table"
        ]
    )

    for element in elements:

        if element.name == "h2":

            text = clean_text(
                element.get_text()
            )

            if text:

                content.append(
                    f"## {text}"
                )

        elif element.name == "h3":

            text = clean_text(
                element.get_text()
            )

            if text:

                content.append(
                    f"### {text}"
                )

        elif element.name == "p":

            # Avoid duplicate text
            # if paragraph is inside list item
            if element.find_parent("li"):
                continue

            text = clean_text(
                element.get_text()
            )

            if text:

                content.append(text)

        elif element.name == "li":

            text = clean_text(
                element.get_text()
            )

            if text:

                content.append(
                    f"- {text}"
                )

        elif element.name == "table":

            table_text = extract_table(
                element
            )

            if table_text:

                content.append(
                    table_text
                )

    return "\n\n".join(content)


# --------------------------------------------------
# EXTRACT ARTICLE
# --------------------------------------------------

def extract_article(url):

    soup = get_soup(url)

    article = soup.select_one(
        "article.ast-article-single"
    )

    if not article:

        raise Exception(
            "Article container not found"
        )

    title_element = article.select_one(
        ".entry-header h1.entry-title"
    )

    author_element = article.select_one(
        ".entry-meta .author-name"
    )

    date_element = article.select_one(
        ".entry-meta .published"
    )

    body = article.select_one(
        ".entry-content.clear"
    )

    if not body:

        raise Exception(
            "Article body not found"
        )

    h2_headings = [

        clean_text(
            heading.get_text()
        )

        for heading in body.select(
            "h2.wp-block-heading"
        )
    ]

    h3_headings = [

        clean_text(
            heading.get_text()
        )

        for heading in body.select(
            "h3.wp-block-heading"
        )
    ]

    tables = []

    for table in body.select(
        "figure.wp-block-table table"
    ):

        tables.append(
            extract_table(table)
        )

    categories = []

    for class_name in article.get(
        "class",
        []
    ):

        if class_name.startswith(
            "category-"
        ):

            categories.append(
                class_name.replace(
                    "category-",
                    ""
                )
            )

    return {

        "Article_Title":
            clean_text(
                title_element.get_text()
                if title_element
                else ""
            ),

        "Author":
            clean_text(
                author_element.get_text()
                if author_element
                else ""
            ),

        "Published_Date":
            clean_text(
                date_element.get_text()
                if date_element
                else ""
            ),

        "Categories":
            ", ".join(
                categories
            ),

        "H2_Headings":
            " | ".join(
                h2_headings
            ),

        "H3_Headings":
            " | ".join(
                h3_headings
            ),

        "Tables":
            "\n\n--- TABLE ---\n\n".join(
                tables
            ),

        "Full_Article":
            extract_article_body(
                body
            )
    }


# --------------------------------------------------
# CHECKPOINT
# --------------------------------------------------

def load_completed_urls():

    if not os.path.exists(
        CHECKPOINT_FILE
    ):

        return set()

    try:

        dataframe = pd.read_csv(
            CHECKPOINT_FILE
        )

        completed = dataframe[
            dataframe["Status"]
            ==
            "OK"
        ]

        return set(
            completed[
                "Blog_URL"
            ].dropna()
        )

    except Exception:

        return set()


def save_checkpoint(row):

    dataframe = pd.DataFrame(
        [row]
    )

    file_exists = os.path.exists(
        CHECKPOINT_FILE
    )

    dataframe.to_csv(

        CHECKPOINT_FILE,

        mode="a",

        header=not file_exists,

        index=False,

        encoding="utf-8-sig"
    )


# --------------------------------------------------
# EXCEL CHARACTER LIMIT
# --------------------------------------------------

def split_long_text(
    text,
    max_length=30000
):

    if not text:
        return [""]

    chunks = []

    while len(text) > max_length:

        split_position = text.rfind(
            "\n\n",
            0,
            max_length
        )

        if split_position == -1:

            split_position = max_length

        chunks.append(
            text[
                :split_position
            ].strip()
        )

        text = text[
            split_position:
        ].strip()

    if text:

        chunks.append(text)

    return chunks


# --------------------------------------------------
# CREATE FINAL EXCEL
# --------------------------------------------------

def create_excel():

    dataframe = pd.read_csv(
        CHECKPOINT_FILE
    )

    dataframe = dataframe.drop_duplicates(

        subset=[
            "Blog_URL"
        ],

        keep="last"
    )

    article_chunks = dataframe[
        "Full_Article"
    ].fillna("").apply(
        split_long_text
    )

    max_chunks = max(
        article_chunks.apply(len)
    )

    for number in range(
        max_chunks
    ):

        dataframe[
            f"Article_Part_{number + 1}"
        ] = article_chunks.apply(

            lambda parts:

            parts[number]

            if number < len(parts)

            else ""
        )

    dataframe.drop(

        columns=[
            "Full_Article"
        ],

        inplace=True
    )

    dataframe.insert(

        0,

        "Sr_No",

        range(
            1,
            len(dataframe) + 1
        )
    )

    dataframe.to_excel(

        OUTPUT_FILE,

        index=False
    )

    format_excel()


# --------------------------------------------------
# FORMAT EXCEL
# --------------------------------------------------

def format_excel():

    workbook = load_workbook(
        OUTPUT_FILE
    )

    sheet = workbook.active

    sheet.title = "Blogs"

    sheet.freeze_panes = "A2"

    sheet.auto_filter.ref = (
        sheet.dimensions
    )

    for cell in sheet[1]:

        cell.font = Font(
            bold=True
        )

    for row in sheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(

                vertical="top",

                wrap_text=True
            )

    for column in sheet.columns:

        letter = get_column_letter(
            column[0].column
        )

        header = column[0].value

        if str(header).startswith(
            "Article_Part_"
        ):

            width = 80

        elif header in [
            "Card_Title",
            "Article_Title",
            "Excerpt",
            "H2_Headings",
            "H3_Headings"
        ]:

            width = 55

        elif header in [
            "Blog_URL",
            "Thumbnail_URL"
        ]:

            width = 60

        elif header == "Tables":

            width = 70

        else:

            width = 20

        sheet.column_dimensions[
            letter
        ].width = width

    workbook.save(
        OUTPUT_FILE
    )


# --------------------------------------------------
# MAIN PROGRAM
# --------------------------------------------------

def main():

    print(
        "\n=============================="
    )

    print(
        "TechnoEdge FULL Blog Scraper"
    )

    print(
        "==============================\n"
    )

    print(
        "Reading first archive page..."
    )

    first_page = get_soup(
        BASE_URL
    )

    total_pages = get_total_pages(
        first_page
    )

    print(
        f"Total archive pages detected: "
        f"{total_pages}"
    )

    all_cards = []

    # -----------------------------------
    # ARCHIVE PAGES
    # -----------------------------------

    for page_number in range(
        1,
        total_pages + 1
    ):

        if page_number == 1:

            soup = first_page

            page_url = BASE_URL

        else:

            page_url = (
                f"{BASE_URL}"
                f"page/"
                f"{page_number}/"
            )

            print(
                f"\nOpening archive page "
                f"{page_number}:"
            )

            print(page_url)

            soup = get_soup(
                page_url
            )

        cards = extract_cards(
            soup,
            page_number
        )

        print(
            f"Page {page_number}: "
            f"{len(cards)} cards found"
        )

        all_cards.extend(
            cards
        )

        time.sleep(
            REQUEST_DELAY
        )

    cards_dataframe = pd.DataFrame(
        all_cards
    )

    cards_dataframe = (
        cards_dataframe
        .drop_duplicates(
            subset=[
                "Blog_URL"
            ]
        )
    )

    total_articles = len(
        cards_dataframe
    )

    print(
        "\n=============================="
    )

    print(
        f"Unique articles found: "
        f"{total_articles}"
    )

    print(
        "==============================\n"
    )

    completed_urls = (
        load_completed_urls()
    )

    # -----------------------------------
    # ARTICLE PAGES
    # -----------------------------------

    for number, (_, card) in enumerate(

        cards_dataframe.iterrows(),

        start=1
    ):

        url = card[
            "Blog_URL"
        ]

        if url in completed_urls:

            print(
                f"[{number}/"
                f"{total_articles}] "
                f"Already completed - SKIPPED"
            )

            continue

        print(
            f"\n[{number}/"
            f"{total_articles}]"
        )

        print(
            card[
                "Card_Title"
            ]
        )

        print(
            url
        )

        row = card.to_dict()

        try:

            article_data = (
                extract_article(
                    url
                )
            )

            row.update(
                article_data
            )

            row[
                "Status"
            ] = "OK"

            row[
                "Error"
            ] = ""

            print(
                "SUCCESS"
            )

        except Exception as error:

            row.update({

                "Article_Title":
                    "",

                "Author":
                    "",

                "Published_Date":
                    "",

                "Categories":
                    "",

                "H2_Headings":
                    "",

                "H3_Headings":
                    "",

                "Tables":
                    "",

                "Full_Article":
                    "",

                "Status":
                    "FAILED",

                "Error":
                    str(error)
            })

            print(
                f"FAILED: {error}"
            )

        save_checkpoint(
            row
        )

        time.sleep(
            REQUEST_DELAY
        )

    # -----------------------------------
    # FINAL EXCEL
    # -----------------------------------

    print(
        "\nCreating Excel file..."
    )

    create_excel()

    print(
        "\n=============================="
    )

    print(
        "SCRAPING COMPLETED"
    )

    print(
        "=============================="
    )

    print(
        f"\nFinal Excel:\n"
        f"{OUTPUT_FILE}"
    )

    print(
        f"\nCheckpoint:\n"
        f"{CHECKPOINT_FILE}"
    )


if __name__ == "__main__":

    main()